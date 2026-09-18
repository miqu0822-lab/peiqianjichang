# -*- coding: utf-8 -*-
import sys, copy
sys.path.insert(0, '.')
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_lib import SRC, find_para, set_text, clone_para, build_table
import content as C

doc = docx.Document(SRC)
body = doc.element.body

# ---- templates ---------------------------------------------------------
tpl_body = find_para(doc, "The results distinguish checkpoint-level geometric fidelity")
tpl_cap  = find_para(doc, "Table 7. Image-evaluation results at high resolution")
tpl_h1   = find_para(doc, "Appendix B Drift Benchmarks")
tpl_h2   = find_para(doc, "4.1 Image Fidelity and Objective Specific Optima")
tpl_h3   = find_para(doc, "3.4.1 HSV Structural Classification")

def P(text):  return clone_para(tpl_body, text)
def CAP(text):return clone_para(tpl_cap, text)
def H1(text): return clone_para(tpl_h1, text)
def H2(text): return clone_para(tpl_h2, text)
def H3(text): return clone_para(tpl_h3, text)

def insert_before(anchor_p, elems):
    for e in elems:
        anchor_p._p.addprevious(e)

def move_before(anchor_p, elems):
    for e in elems:
        anchor_p._p.addprevious(e)

# =======================================================================
# A. Abstract
set_text(find_para(doc, "Image fidelity alone does not establish"), C.ABSTRACT)

# B. Highlights
hl_keys = ["A factorial GAN comparison separates",
           "Geometric metrics identify different configuration leaders",
           "Nineteen ETABS models connect generated layouts",
           "Numerical quality-control success does not establish drift acceptability.",
           "Image ranking and mechanical assessment serve distinct"]
for key, new in zip(hl_keys, C.HIGHLIGHTS):
    set_text(find_para(doc, key), new)

# C. Introduction consistency
set_text(find_para(doc, "Twelve single-factor contrasts describe checkpoint performance"), C.INTRO_P17)

# D. Section 3.4.4 cross-reference  Table 10 -> Table 13
p = find_para(doc, "The adopted ScoreIoU gives equal weight")
set_text(p, p.text.replace("Table 10 compares", "Table 13 compares"))

# E. Section 4.1
set_text(find_para(doc, "4.1 Image Fidelity and Objective Specific Optima"), "4.1 Checkpoint-Level Image Fidelity")
set_text(find_para(doc, "At HR, A7 had the highest observed ScoreIoU"), C.P105_NEW)

# =======================================================================
# F. New Section 4.2 (inserted before the old 4.2 heading)
anchor_42 = find_para(doc, "4.2 Metric Decomposition and Parameter Cost")

W7  = [2190, 1080, 1180, 1120, 1080, 1180, 1196]           # Table 9  (7 cols)
W4  = [2400, 2000, 2200, 2426]                              # Table 10 (4 cols)
W11 = [1700, 1420, 1080, 1420, 1080, 1000, 1326]            # Table 11 (7 cols)
W12 = [1500, 1000, 980, 1180, 900, 1180, 1030, 1256]        # Table 12 (8 cols)

blocks = [
    H2(C.S42_HEAD),
    P(C.S42_P1),
    CAP(C.T9_CAP), build_table(C.T9_ROWS, W7), P(C.T9_NOTE),
    P(C.S42_P2),
    P(C.S42_P3),
    CAP(C.T10_CAP), build_table(C.T10_ROWS, W4), P(C.T10_NOTE),
    P(C.S42_P4),
    P(C.S42_P5),
    CAP(C.T11_CAP), build_table(C.T11_ROWS, W11), P(C.T11_NOTE),
]
insert_before(anchor_42, blocks)

# Figure 9 picture + caption
pic_par = doc.add_paragraph()
pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_par.paragraph_format.first_line_indent = 0
pic_par.add_run().add_picture('Figure9_factorial.png', width=Inches(6.27))
insert_before(anchor_42, [pic_par._p, CAP(C.FIG9_CAP)])

# 4.2.1
insert_before(anchor_42, [H3(C.S421_HEAD), P(C.S421_P1), P(C.S421_P2), P(C.S421_P3)])

# =======================================================================
# G. old 4.2 -> 4.3
set_text(anchor_42, C.S43_HEAD)

# H. new material at the end of 4.3, after the parameter-count paragraph
anchor_43_end = find_para(doc, "4.3 Descriptive Component Contrasts")
insert_before(anchor_43_end, [
    P(C.S43_PA),
    CAP(C.T12_CAP), build_table(C.T12_ROWS, W12), P(C.T12_NOTE),
    P(C.S43_PB),
    P(C.S43_PC),
    P(C.S43_T13_LEAD),
])

# I/J. move the aggregation table (old Table 10 -> 13) and its discussion into 4.3,
#      then relocate the 12-contrast table to Appendix C.
cap_t10 = find_para(doc, "Table 10. Sensitivity of architecture rankings")
set_text(cap_t10, cap_t10.text.replace("Table 10.", "Table 13."))
tbl_t10 = cap_t10._p.getnext()
assert tbl_t10.tag.endswith('}tbl'), "expected the aggregation table after its caption"
p122 = find_para(doc, "The adopted aggregation ranked A7 first at HR")
move_before(anchor_43_end, [cap_t10._p, tbl_t10, p122._p])

# remove the old 4.3 lead-in and heading
p117 = find_para(doc, "Table 9 reports 12 relative ScoreIoU contrasts")
p117._p.getparent().remove(p117._p)

cap_t9_old = find_para(doc, "Table 9. Descriptive relative ScoreIoU changes")
set_text(cap_t9_old, C.T17_CAP)
tbl_t9_old = cap_t9_old._p.getnext()
assert tbl_t9_old.tag.endswith('}tbl')
p123 = find_para(doc, "All four HR fusion contrasts favored cross-attention")
p124 = find_para(doc, "For unregularized HR models, replacing DCE with CLIP")
p125 = find_para(doc, "Adding SN+EMA changed HR ScoreIoU")
p126 = find_para(doc, "Geometric objectives identify candidate configurations")
set_text(p126, C.APXC_CLOSE)
apxc_items = [cap_t9_old._p, tbl_t9_old, p123._p, p124._p, p125._p, p126._p]

old_43_head = anchor_43_end            # "4.3 Descriptive Component Contrasts"

# =======================================================================
# K. new 4.4 robustness section, replacing the now-empty old 4.3 heading slot
insert_before(old_43_head, [H2(C.S44_HEAD), P(C.S44_P1), P(C.S44_P2)])
old_43_head._p.getparent().remove(old_43_head._p)

# L. old 4.4 -> 4.5, ETABS cross-references renumbered
set_text(find_para(doc, "4.4 Numerical Quality Control and Elastic Response"),
         "4.5 Numerical Quality Control and Elastic Response")
p = find_para(doc, "The 16 matched 8dH1H2 models contained")
set_text(p, p.text.replace("(Figure 9)", "(Figure 10)"))
p = find_para(doc, "Figure 9. Geometry and response of the reconstructed P01 models")
set_text(p, p.text.replace("Figure 9.", "Figure 10.", 1))
p = find_para(doc, "Table 11. Reconstructed geometry and ETABS 22 responses")
set_text(p, p.text.replace("Table 11.", "Table 14.", 1))
p = find_para(doc, "Table 12. Condition-specific responses of the three A7-HR")
set_text(p, p.text.replace("Table 12.", "Table 15.", 1))

# M. old 4.5 -> 4.6 and the two mechanistic paragraphs
set_text(find_para(doc, "4.5 Mechanistic Interpretation of the Observed Results"),
         "4.6 Mechanistic Interpretation of the Observed Results")
set_text(find_para(doc, "The SN+EMA treatment behaves as a stability"), C.S46_SNEMA)
set_text(find_para(doc, "The fact that the fully augmented A5"), C.S46_A5)

# N. Conclusion
set_text(find_para(doc, "Three findings emerge from the factorial comparison"), C.C_P2)
p = find_para(doc, "The reconstruction and elastic-assessment stage confirms")
old = p.text
tail = old.split("only loosely coupled.", 1)[1]
set_text(p, C.C_P3_FIRST + tail)
set_text(find_para(doc, "Taken together, the results argue that"), C.C_P4)

# O. Appendix B cross-references
p = find_para(doc, "adds diagnostics calculated from the existing ETABS exports")
set_text(p, p.text.replace("Table 13 adds", "Table 16 adds").replace("values in Table 11", "values in Table 14"))
p = find_para(doc, "Table 13. Drift utilization and directional base-shear")
set_text(p, p.text.replace("Table 13.", "Table 16.", 1))

# P. Appendix C
refs = find_para(doc, "References")
insert_before(refs, [H1(C.APXC_HEAD), P(C.APXC_LEAD)])
move_before(refs, apxc_items)

out = '/home/user/peiqianjichang/revision/manuscript_revised.docx'
doc.save(out)
print("saved", out)
